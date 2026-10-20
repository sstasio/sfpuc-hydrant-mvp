"""Build the Survey123 XLSForm for hydrant inspections from the schema in build_hydrant_gdb.py.

    python build_xlsform.py            -> survey123/HydrantInspection.xlsx

Needs openpyxl only (no arcpy). Field names and choice codes come from the same
definitions the geodatabase script uses, so the form matches the published table.
"""
import os
import re

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill

import build_hydrant_gdb as schema

OUT = os.path.join(schema.HERE, "survey123", "HydrantInspection.xlsx")

SURVEY_COLS = ["type", "name", "label", "hint", "required", "relevant", "constraint",
               "constraint_message", "calculation", "choice_filter", "appearance", "default",
               "readonly", "bind::esri:fieldType", "bind::esri:fieldLength"]


def q(type_, name, label="", **kw):
    row = {"type": type_, "name": name, "label": label}
    row.update(kw)
    return row


def yes_if_selected(list_question, code):
    return f"if(selected(${{{list_question}}}, '{code}'), 'Yes', 'No')"


def survey_rows():
    psi = dict(constraint=". >= 0 and . <= 300", constraint_message="Enter 0 to 300 psi")
    flags = [f"count-selected(${{{lst}}}) > 0" for lst in ("markings", "issues")]
    needs_followup = (
        "if(${FoundAtLoc} = 'No' or ${Operable} != 'Yes' or ${LowPressure} = 'Yes'"
        " or number(${CapsMissing}) > 0 or (string-length(${ShutoffValve}) > 0 and ${ShutoffValve} != 'OK')"
        f" or ${{Obstructed}} = 'Yes' or {' or '.join(flags)}, 'Yes', 'No')"
    )
    rows = [
        # Prefilled from the Field Maps pop-up link (field:hydrantguid={GlobalID} etc.)
        q("hidden", "HydrantGUID", **{"bind::esri:fieldType": "esriFieldTypeGUID"}),
        q("begin group", "hydrant", "Hydrant", appearance="field-list"),
        q("text", "GISID", "GIS ID", required="yes", readonly="yes",
          hint="Filled in from the map. If blank, open this form from the hydrant's pop-up."),
        q("text", "Location", "Address/Location", readonly="yes"),
        q("hidden", "Lat", **{"bind::esri:fieldType": "esriFieldTypeDouble"}),
        q("hidden", "Lon", **{"bind::esri:fieldType": "esriFieldTypeDouble"}),
        q("end group", "hydrant"),

        q("begin group", "crew", "Crew", appearance="field-list"),
        q("dateTime", "InspDate", "Inspection Date", required="yes", default="now()"),
        q("select_one EngineCompany", "EngineCo", "Engine Company", required="yes", appearance="autocomplete"),
        q("select_one Inspector", "Inspector", "Inspector", required="yes", appearance="autocomplete",
          choice_filter="engineco=${EngineCo} or name='Other'",
          hint="Names for the selected engine company. Pick Other and add your name in Remarks if missing."),
        q("end group", "crew"),

        q("select_one YesNo", "FoundAtLoc", "Hydrant found at this location?", required="yes",
          appearance="horizontal"),

        q("begin group", "condition", "Condition", relevant="${FoundAtLoc} = 'Yes'", appearance="field-list"),
        q("select_one Operable", "Operable", "Operable", required="yes", appearance="horizontal"),
        q("select_one YesNo", "Corrosion", "Corrosion", required="yes", appearance="horizontal"),
        q("integer", "Gauge", "Gauge Pressure (psi)", hint="Static pressure read on the gauge. Leave blank if not read.", **psi),
        q("integer", "Spindle", "Spindle Pressure (psi)", hint="Pressure stamped on the spindle. Leave blank if not legible.", **psi),
        q("select_one YesNo", "LowPressure", "Low Pressure", readonly="yes", appearance="horizontal",
          calculation="if(string-length(${Gauge}) > 0 and string-length(${Spindle}) > 0"
                      " and number(${Gauge}) < number(${Spindle}), 'Yes', 'No')"),
        q("integer", "CapsMissing", "Caps Missing", default="0",
          constraint=". >= 0 and . <= 3", constraint_message="Enter 0 to 3"),
        q("select_one ShutoffValve", "ShutoffValve", "Shutoff Valve"),
        q("select_one YesNo", "Obstructed", "Obstructed", appearance="horizontal"),
        q("select_multiple markings", "markings", "Markings missing or unreadable",
          **{"bind::esri:fieldType": "null"}),
        q("select_multiple issues", "issues", "Other issues", **{"bind::esri:fieldType": "null"}),
    ]
    # Split the two checklists into the separate Yes/No fields the table stores.
    for lst, flags_ in (("markings", schema.MARKING_FLAGS), ("issues", schema.ISSUE_FLAGS)):
        rows += [q("hidden", n, calculation=yes_if_selected(lst, n),
                   **{"bind::esri:fieldType": "esriFieldTypeString", "bind::esri:fieldLength": "3"})
                 for n, _ in flags_]
    rows += [
        q("end group", "condition"),
        q("select_one YesNo", "NeedsFollowup", "Needs SFPUC Follow-up", appearance="horizontal",
          calculation=needs_followup, hint="Set automatically when an issue is flagged. You can change it."),
        q("text", "Remarks", "Remarks", appearance="multiline", **{"bind::esri:fieldLength": "1000"}),
        q("image", "photo1", "Photo", required="${NeedsFollowup} = 'Yes'",
          hint="Required when follow-up is needed."),
        q("image", "photo2", "Second photo (optional)"),
    ]
    return rows


def portal_names(rows):
    """Lower-case question names and ${} references.

    The portal lower-cases every field name when the layer is published, and
    Survey123 matches questions to fields by name.
    """
    for r in rows:
        r["name"] = r["name"].lower()
        for k in ("relevant", "calculation", "choice_filter", "required"):
            if k in r:
                r[k] = re.sub(r"\$\{(\w+)\}", lambda m: "${%s}" % m.group(1).lower(), r[k])
    return rows


def choice_rows():
    domains = schema.coded_domains()
    rows = []
    for list_name in ("YesNo", "EngineCompany", "Operable", "ShutoffValve"):
        rows += [{"list_name": list_name, "name": c, "label": l} for c, l in domains[list_name][1]]
    # An inspector seen under more than one company gets one row per company
    # (allowed by allow_choice_duplicates in settings).
    for c, n in sorted(schema.load_inspectors(), key=lambda r: (r[1], r[0])):
        rows.append({"list_name": "Inspector", "name": schema.inspector_code(n),
                     "label": n.replace(",", ", "), "engineco": c})
    rows.append({"list_name": "Inspector", "name": "Other", "label": "Other (type in Remarks)"})
    rows += [{"list_name": "markings", "name": n, "label": l} for n, l in schema.MARKING_FLAGS]
    rows += [{"list_name": "issues", "name": n, "label": l} for n, l in schema.ISSUE_FLAGS]
    return rows


def write_sheet(ws, cols, rows):
    ws.append(cols)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="004A80")
    for r in rows:
        ws.append([r.get(c, "") for c in cols])
    for i, c in enumerate(cols, 1):
        ws.column_dimensions[ws.cell(1, i).column_letter].width = max(12, min(60, len(c) + 4))
    ws.freeze_panes = "A2"


def main():
    wb = Workbook()
    write_sheet(wb.active, SURVEY_COLS, portal_names(survey_rows()))
    wb.active.title = "survey"
    write_sheet(wb.create_sheet("choices"), ["list_name", "name", "label", "engineco"], choice_rows())
    write_sheet(wb.create_sheet("settings"),
                ["form_title", "form_id", "instance_name", "submission_url", "version",
                 "allow_choice_duplicates"],
                [{"form_title": "Hydrant Inspection", "form_id": schema.INSPECTIONS,
                  "instance_name": "concat(${gisid}, ' ', format-date(${inspdate}, '%Y-%m-%d'))",
                  "submission_url": "", "version": "1.0",
                  "allow_choice_duplicates": "yes"}])
    os.makedirs(os.path.dirname(OUT), exist_ok=True)
    wb.save(OUT)
    print(f"Wrote {OUT}")


if __name__ == "__main__":
    main()
