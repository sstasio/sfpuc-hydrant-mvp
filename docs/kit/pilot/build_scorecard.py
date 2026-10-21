"""Build pilot-scorecard.xlsx from answer_key.csv (needs openpyxl).

    python build_scorecard.py            # writes pilot-scorecard.xlsx next to this file
    python build_scorecard.py --testers 10
"""
import argparse
import csv
import os

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

HERE = os.path.dirname(os.path.abspath(__file__))
FM, S123 = "Field Maps", "Survey123"
ORDERS = [(FM, "A"), (S123, "A"), (FM, "B"), (S123, "B")]

# SFPUC app theme (sfpuc-design-system)
HEADER = PatternFill("solid", fgColor="004A80")
INPUT = PatternFill("solid", fgColor="FFF8E1")
CALC = PatternFill("solid", fgColor="F4F6F8")
PASS_FILL = PatternFill("solid", fgColor="E3EED3")
FAIL_FILL = PatternFill("solid", fgColor="F8D7D5")
WHITE_BOLD = Font(name="Arial", bold=True, color="FFFFFF")
BOLD = Font(name="Arial", bold=True)
BASE = Font(name="Arial")
TITLE = Font(name="Arial", bold=True, size=14, color="004A80")
THIN = Side(style="thin", color="C9D1D9")
BOX = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
WRAP = Alignment(wrap_text=True, vertical="top")

TASK_NAMES = {"T1": "Routine check", "T2": "Low pressure, missing cap", "T3": "Several problems",
              "T4": "Hydrant not there", "T5": "No signal (offline)"}

TEAM_CRITERIA = [
    ("Account and licensing cost", "No added licenses for 44 engine companies", 15),
    ("Inspections tied to the right hydrant", "Every record lands under its hydrant without cleanup", 25),
    ("Offline", "Works with no signal, map included, syncs without fuss", 20),
    ("Upkeep", "Form and inspector list are quick to change and roll out", 15),
    ("Dashboard and annual report", "Dashboard and FY report come straight from the data", 10),
    ("Later Maximo link (FME Flow)", "Follow-up flags can become work requests with no rework", 10),
    ("Security and sharing", "Fits SFPUC portal policy with no exceptions", 5),
]


def header(ws, row, labels, widths=None):
    for i, label in enumerate(labels, 1):
        c = ws.cell(row=row, column=i, value=label)
        c.fill, c.font, c.border = HEADER, WHITE_BOLD, BOX
        c.alignment = Alignment(wrap_text=True, vertical="center")
        if widths:
            ws.column_dimensions[get_column_letter(i)].width = widths[i - 1]
    ws.row_dimensions[row].height = 32
    ws.freeze_panes = f"A{row + 1}"  # a string, so no empty cell is created in that row


def style_rows(ws, first, last, ncols, input_cols=()):
    for r in range(first, last + 1):
        for c in range(1, ncols + 1):
            cell = ws.cell(row=r, column=c)
            cell.border, cell.font = BOX, BASE
            cell.alignment = Alignment(vertical="top", wrap_text=True)
            cell.fill = INPUT if c in input_cols else CALC


def add_list(ws, values, ref):
    dv = DataValidation(type="list", formula1='"' + ",".join(values) + '"', allow_blank=True)
    dv.add(ref)
    ws.add_data_validation(dv)


def add_whole(ws, lo, hi, ref):
    dv = DataValidation(type="whole", operator="between", formula1=str(lo), formula2=str(hi), allow_blank=True,
                        error=f"Enter a whole number from {lo} to {hi}", showErrorMessage=True)
    dv.add(ref)
    ws.add_data_validation(dv)


def landscape(ws):
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    ws.print_title_rows = "1:1"


def build(path, n_testers):
    key = list(csv.DictReader(open(os.path.join(HERE, "answer_key.csv"), newline="", encoding="utf-8")))
    gisid = {(r["Task"], r["Set"]): r["GISID"] for r in key}

    wb = Workbook()

    # ReadMe
    ws = wb.active
    ws.title = "ReadMe"
    ws.column_dimensions["A"].width = 110
    lines = [
        ("Hydrant inspection pilot scorecard", TITLE),
        ("Field Maps vs Survey123. Run the sessions from PILOT-SCRIPT.md.", BASE),
        ("", BASE),
        ("Yellow cells are for typing. Grey cells are filled in for you or calculated.", BOLD),
        ("", BASE),
        ("Testers: one row per tester. Fill in company, phone and account before their session.", BASE),
        ("Tasks: the observer logs time, result and ease for every task, in the order the rows are listed.", BASE),
        ("After-app: the four questions after each app (1 strongly disagree, 5 strongly agree).", BASE),
        ("Wrap-up: which app the tester would pick, and why.", BASE),
        ("Data check: paste the output of check_pilot_records.py here (columns already match).", BASE),
        ("Team criteria: the team scores each app 1 to 5. Agree on the weights before scoring.", BASE),
        ("Summary: adds everything up and shows the must-pass checks and the combined score.", BASE),
        ("Observer sheet: print one per tester if you'd rather take notes on paper.", BASE),
        ("Answer key: the right values for every task card.", BASE),
        ("", BASE),
        ("Combined score = crew score and team score weighted equally, each out of 100.", BASE),
        ("Crew score = 40% tasks done alone + 30% average ease + 30% after-app agreement (tasks 1 to 5 only, "
         "so both apps are scored on the same tasks).", BASE),
        ("Team score = weighted Team criteria, scaled to 100.", BASE),
    ]
    for i, (text, font) in enumerate(lines, 1):
        c = ws.cell(row=i, column=1, value=text)
        c.font, c.alignment = font, Alignment(wrap_text=True)

    # Testers
    ws = wb.create_sheet("Testers")
    header(ws, 1, ["Pilot ID", "Engine company", "Phone", "Has portal account", "First app", "First set",
                   "Session date", "Notes"], [10, 16, 12, 18, 14, 10, 14, 50])
    for i in range(1, n_testers + 1):
        app, s = ORDERS[(i - 1) % 4]
        ws.append([f"P{i}", None, None, None, app, s, None, None])
    last = n_testers + 1
    style_rows(ws, 2, last, 8, input_cols=(2, 3, 4, 7, 8))
    add_list(ws, [f"E{n:02d}" for n in range(1, 52)], f"B2:B{last}")
    add_list(ws, ["iPhone", "Android"], f"C2:C{last}")
    add_list(ws, ["Yes", "No"], f"D2:D{last}")
    for r in range(2, last + 1):
        ws.cell(row=r, column=7).number_format = "yyyy-mm-dd"

    # Tasks
    ws = wb.create_sheet("Tasks")
    header(ws, 1, ["Pilot ID", "Block", "App", "Set", "Task", "Task name", "GIS ID", "Minutes",
                   "Result", "Ease (1-7)", "Photo taken", "Notes"],
           [9, 8, 12, 6, 6, 26, 8, 9, 12, 9, 10, 60])
    r = 2
    for i in range(1, n_testers + 1):
        first_app, first_set = ORDERS[(i - 1) % 4]
        blocks = [(1, first_app, first_set),
                  (2, S123 if first_app == FM else FM, "B" if first_set == "A" else "A")]
        for block, app, s in blocks:
            for t in ["T1", "T2", "T3", "T4", "T5"]:
                ws.append([f"P{i}", block, app, s, t, TASK_NAMES[t], gisid[(t, s)]])
                r += 1
    last = r - 1
    style_rows(ws, 2, last, 12, input_cols=(8, 9, 10, 11, 12))
    add_list(ws, ["Alone", "With help", "Failed", "Skipped"], f"I2:I{last}")
    add_whole(ws, 1, 7, f"J2:J{last}")
    add_list(ws, ["Yes", "No", "Not needed"], f"K2:K{last}")
    dv = DataValidation(type="decimal", operator="between", formula1="0", formula2="60", allow_blank=True)
    dv.add(f"H2:H{last}")
    ws.add_data_validation(dv)
    for rr in range(2, last + 1):
        ws.cell(row=rr, column=8).number_format = "0.0"
    landscape(ws)
    tasks_last = last

    # After-app
    ws = wb.create_sheet("After-app")
    header(ws, 1, ["Pilot ID", "App", "Q1 Could do it without help", "Q2 Quick enough for a shift",
                   "Q3 Easy one-handed and outdoors", "Q4 Trust it saved correctly", "Average",
                   "Most frustrating part"], [9, 12, 14, 14, 14, 14, 9, 60])
    r = 2
    for i in range(1, n_testers + 1):
        first_app, _ = ORDERS[(i - 1) % 4]
        for app in [first_app, S123 if first_app == FM else FM]:
            ws.append([f"P{i}", app, None, None, None, None,
                       f'=IF(COUNT(C{r}:F{r})=0,"",AVERAGE(C{r}:F{r}))'])
            r += 1
    after_last = r - 1
    style_rows(ws, 2, after_last, 8, input_cols=(3, 4, 5, 6, 8))
    add_whole(ws, 1, 5, f"C2:F{after_last}")
    for rr in range(2, after_last + 1):
        ws.cell(row=rr, column=7).number_format = "0.0"

    # Wrap-up
    ws = wb.create_sheet("Wrap-up")
    header(ws, 1, ["Pilot ID", "Would pick", "Why", "Anything neither app handled"], [9, 16, 60, 50])
    for i in range(1, n_testers + 1):
        ws.append([f"P{i}"])
    style_rows(ws, 2, n_testers + 1, 4, input_cols=(2, 3, 4))
    add_list(ws, [FM, S123, "No preference"], f"B2:B{n_testers + 1}")

    # Data check
    ws = wb.create_sheet("Data check")
    header(ws, 1, ["PilotID", "Task", "Set", "App", "GISID", "Linked", "WrongCount", "WrongFields",
                   "ObjectID", "Note"], [9, 6, 6, 12, 8, 13, 11, 60, 10, 40])
    data_last = 1 + n_testers * 11 + 20
    style_rows(ws, 2, data_last, 10, input_cols=range(1, 11))
    landscape(ws)

    # Team criteria
    ws = wb.create_sheet("Team criteria")
    header(ws, 1, ["Criterion", "What a 5 looks like", "Weight", "Field Maps (1-5)", "Survey123 (1-5)",
                   "Field Maps weighted", "Survey123 weighted", "Notes"], [34, 50, 9, 12, 12, 12, 12, 50])
    for name, five, weight in TEAM_CRITERIA:
        r = ws.max_row + 1
        ws.append([name, five, weight, None, None,
                   f'=IF(D{r}="","",C{r}*D{r})', f'=IF(E{r}="","",C{r}*E{r})'])
    crit_last = ws.max_row
    tot = crit_last + 1
    ws.append(["Total", None, f"=SUM(C2:C{crit_last})", None, None,
               f"=SUM(F2:F{crit_last})", f"=SUM(G2:G{crit_last})"])
    style_rows(ws, 2, tot, 8, input_cols=(3, 4, 5, 8))
    for c in range(1, 9):
        ws.cell(row=tot, column=c).font = BOLD
    add_whole(ws, 1, 5, f"D2:E{crit_last}")

    # Summary
    ws = wb.create_sheet("Summary", 1)
    ws.column_dimensions["A"].width = 52
    ws.column_dimensions["B"].width = 16
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 50
    ws["A1"] = "Pilot summary"
    ws["A1"].font = TITLE
    header(ws, 3, ["Measure", FM, S123, "How it's worked out"])
    ws.freeze_panes = None

    col = lambda letter, sheet="Tasks", last=tasks_last: f"{sheet}!${letter}$2:${letter}${last}"
    tapp, ttask, tmin, tres, tease = col("C"), col("E"), col("H"), col("I"), col("J")
    aapp, aavg = col("B", "'After-app'", after_last), col("G", "'After-app'", after_last)
    dapp, dtask, dlink, dwrong = (col("D", "'Data check'", data_last), col("B", "'Data check'", data_last),
                                  col("F", "'Data check'", data_last), col("G", "'Data check'", data_last))
    pick = col("B", "'Wrap-up'", n_testers + 1)

    rows = [
        ("Crew results (tasks 1 to 5)", None, None),
        ("Tasks logged",
         f'=COUNTIFS({tapp},{{APP}},{tres},"<>")', "Tasks with a Result filled in"),
        ("Done alone",
         f'=IFERROR(COUNTIFS({tapp},{{APP}},{tres},"Alone")/[[Tasks logged]],"")', "Share of logged tasks"),
        ("Done with help",
         f'=IFERROR(COUNTIFS({tapp},{{APP}},{tres},"With help")/[[Tasks logged]],"")', ""),
        ("Failed or skipped",
         f'=IFERROR((COUNTIFS({tapp},{{APP}},{tres},"Failed")'
         f'+COUNTIFS({tapp},{{APP}},{tres},"Skipped"))/[[Tasks logged]],"")', ""),
        ("Average minutes per task",
         f'=IFERROR(AVERAGEIFS({tmin},{tapp},{{APP}}),"")', ""),
        ("Average ease (1 to 7)",
         f'=IFERROR(AVERAGEIFS({tease},{tapp},{{APP}}),"")', "Asked after every task"),
        ("Offline task synced",
         f'=IFERROR((COUNTIFS({tapp},{{APP}},{ttask},"T5",{tres},"Alone")'
         f'+COUNTIFS({tapp},{{APP}},{ttask},"T5",{tres},"With help"))'
         f'/COUNTIFS({tapp},{{APP}},{ttask},"T5",{tres},"<>"),"")', "Task 5 done, alone or with help"),
        ("After-app agreement (1 to 5)", f'=IFERROR(AVERAGEIFS({aavg},{aapp},{{APP}}),"")', "Average of Q1 to Q4"),
        ("Testers who would pick it", f'=COUNTIF({pick},{{APP}})', "From Wrap-up"),
        ("Data quality (tasks 1 to 5)", None, None),
        ("Records checked",
         f'=COUNTIFS({dapp},{{APP}},{dwrong},">=0")', "From Data check"),
        ("Tied to the right hydrant",
         f'=IFERROR(COUNTIFS({dapp},{{APP}},{dlink},"Yes")/[[Records checked]],"")', ""),
        ("Records with no mistakes",
         f'=IFERROR(COUNTIFS({dapp},{{APP}},{dwrong},0)/[[Records checked]],"")', ""),
        ("Average wrong fields per record",
         f'=IFERROR(AVERAGEIFS({dwrong},{dapp},{{APP}}),"")', ""),
        ("Must pass", None, None),
        ("80% or more done alone", '=IF([[Done alone]]="","",IF([[Done alone]]>=0.8,"PASS","FAIL"))', ""),
        ("95% or more tied to the right hydrant", '=IF([[Tied to the right hydrant]]="","",IF([[Tied to the right hydrant]]>=0.95,"PASS","FAIL"))', ""),
        ("Offline synced for every tester", '=IF([[Offline task synced]]="","",IF([[Offline task synced]]=1,"PASS","FAIL"))', ""),
        ("Scores (out of 100)", None, None),
        ("Crew score",
         '=IFERROR(100*(0.4*[[Done alone]]+0.3*([[Average ease (1 to 7)]]-1)/6+0.3*([[After-app agreement (1 to 5)]]-1)/4),"")',
         "40% done alone, 30% ease, 30% after-app"),
        ("Team score",
         "=IF('Team criteria'!{TEAMCOL}" + f"{tot}=0,\"\",100*'Team criteria'!{{TEAMCOL}}{tot}/(5*'Team criteria'!$C${tot}))",
         "Weighted Team criteria"),
        ("Combined score", '=IFERROR(([[Crew score]]+[[Team score]])/2,"")', "Crew and team weighted equally"),
        ("Gap between apps", None, None),
    ]
    # [[Label]] in a formula is replaced with that Summary row's cell in the same app column.
    start = 4
    lab = {label: start + i for i, (label, _, _) in enumerate(rows)}
    for i, (label, fm, s123) in enumerate(rows):
        r = start + i
        ws.cell(row=r, column=1, value=label)
        is_section = fm is None and s123 is None
        if is_section:
            ws.cell(row=r, column=1).font = BOLD
            ws.cell(row=r, column=1).fill = PatternFill("solid", fgColor="D9E6F2")
            for c in range(2, 5):
                ws.cell(row=r, column=c).fill = PatternFill("solid", fgColor="D9E6F2")
            continue
        if fm is not None and fm.startswith("="):
            for c, letter, app in ((2, "B", "B$3"), (3, "C", "C$3")):
                f = fm.replace("{APP}", app)
                f = f.replace("{TEAMCOL}", "F$" if letter == "B" else "G$")
                for name, ref in lab.items():
                    f = f.replace(f"[[{name}]]", f"{letter}{ref}")
                ws.cell(row=r, column=c, value=f)
            ws.cell(row=r, column=4, value=s123)
        else:
            ws.cell(row=r, column=3, value=s123)
            ws.cell(row=r, column=2, value="n/a")
        for c in range(1, 5):
            cell = ws.cell(row=r, column=c)
            cell.border = BOX
            cell.font = cell.font if cell.font.bold else BASE
            cell.alignment = Alignment(vertical="top", wrap_text=True)

    pct = ["Done alone", "Done with help", "Failed or skipped", "Offline task synced",
           "Tied to the right hydrant", "Records with no mistakes"]
    for label in pct:
        for c in (2, 3):
            ws.cell(row=lab[label], column=c).number_format = "0%"
    for label in ["Average minutes per task", "Average ease (1 to 7)", "After-app agreement (1 to 5)",
                  "Average wrong fields per record"]:
        for c in (2, 3):
            ws.cell(row=lab[label], column=c).number_format = "0.0"
    for label in ["Crew score", "Team score", "Combined score"]:
        for c in (2, 3):
            ws.cell(row=lab[label], column=c).number_format = "0"
            ws.cell(row=lab[label], column=c).font = BOLD

    g = lab["Gap between apps"]
    comb = lab["Combined score"]
    ws.cell(row=g, column=1, value="Closer than 10%? (then tester preference decides, see the pilot page)")
    ws.cell(row=g, column=1).font = BASE
    ws.cell(row=g, column=2, value=f'=IFERROR(IF(ABS(B{comb}-C{comb})<0.1*MAX(B{comb},C{comb}),"Yes","No"),"")')
    ws.merge_cells(start_row=g, start_column=2, end_row=g, end_column=3)
    ws.cell(row=g, column=4, value="Difference under 10% of the higher score")
    for c in range(1, 5):
        ws.cell(row=g, column=c).border = BOX
        ws.cell(row=g, column=c).alignment = Alignment(vertical="top", wrap_text=True)

    gates = f"B{lab['80% or more done alone']}:C{lab['Offline synced for every tester']}"
    ws.conditional_formatting.add(gates, CellIsRule(operator="equal", formula=['"PASS"'], fill=PASS_FILL))
    ws.conditional_formatting.add(gates, CellIsRule(operator="equal", formula=['"FAIL"'], fill=FAIL_FILL))

    # Observer sheet (printable)
    ws = wb.create_sheet("Observer sheet")
    ws["A1"] = "Observer sheet   Pilot ID: ______   Date: ________   App: ______________   Set: ___"
    ws["A1"].font = BOLD
    header(ws, 3, ["Task", "Start", "Stop", "Minutes", "Alone / Help / Failed", "Ease 1-7", "Photo",
                   "Where they got stuck, what they said"], [8, 9, 9, 9, 14, 9, 8, 70])
    ws.freeze_panes = None
    for t in ["T1", "T2", "T3", "T4", "T5"]:
        ws.append([t])
        ws.row_dimensions[ws.max_row].height = 60
    style_rows(ws, 4, ws.max_row, 8, input_cols=range(2, 9))
    r = ws.max_row + 2
    ws.cell(row=r, column=1, value="After-app (1-5):  Q1 without help ___   Q2 quick enough ___   "
            "Q3 one-handed/outdoors ___   Q4 trust it saved ___").font = BASE
    ws.cell(row=r + 2, column=1, value="Most frustrating part:").font = BASE
    landscape(ws)
    ws.print_title_rows = None

    # Answer key
    ws = wb.create_sheet("Answer key")
    cols = list(key[0].keys())
    header(ws, 1, cols, [7, 5, 8] + [11] * (len(cols) - 3))
    for k in key:
        ws.append([int(v) if v.isdigit() and c not in ("GISID",) else v for c, v in k.items()])
    style_rows(ws, 2, len(key) + 1, len(cols))
    landscape(ws)

    wb.save(path)
    return path


if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--testers", type=int, default=8)
    ap.add_argument("--out", default=os.path.join(HERE, "pilot-scorecard.xlsx"))
    a = ap.parse_args()
    print("Wrote", build(a.out, a.testers))
